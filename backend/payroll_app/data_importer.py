"""已确认 Schema 的正式 Excel/CSV 导入服务。

该模块与 :mod:`template_inspector` 分工明确：结构预览只统计元数据，
本模块才读取业务行、执行映射转换并在 SQLite 事务中写入。
"""

from __future__ import annotations

import csv
import json
import re
import sqlite3
import uuid
from decimal import Decimal, InvalidOperation
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Iterator, Optional

from .organization_db import OrganizationDatabase
from .template_inspector import TemplateInspector


class DataImportError(ValueError):
    """用户文件不满足已确认结构或转换规则时抛出。"""


class DataImporter:
    """将一个已确认的数据表与一个 Excel Sheet 建立一次性导入批次。"""

    def __init__(self, database: OrganizationDatabase, inspector: Optional[TemplateInspector] = None) -> None:
        self.database = database
        self.inspector = inspector or TemplateInspector()

    def inspect_file(self, path: Path, *, page_name: str, sheet_name: Optional[str] = None) -> dict[str, Any]:
        """仅限行读取文件结构，不读取或返回完整业务数据。"""
        source = Path(path)
        if not source.is_file():
            raise DataImportError(f"找不到导入文件: {source}")
        page, fields = self._load_page(page_name)
        expected_headers = [field["excel_name"] for field in fields if field["selected"]]
        if source.suffix.lower() == ".csv":
            sheet_names = ["CSV"]
        else:
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise DataImportError("读取 Excel 需要安装 openpyxl") from exc
            book = load_workbook(source, read_only=True, data_only=False)
            try:
                sheet_names = list(book.sheetnames)
            finally:
                book.close()
        selected_sheet = sheet_name or (sheet_names[0] if len(sheet_names) == 1 else None)
        result: dict[str, Any] = {
            "source_file": source.name,
            "sheet_names": sheet_names,
            "selected_sheet": selected_sheet,
            "required_sheet": page["selected_sheet"] or None,
            "schema_hash": self.database.schema_fingerprint(page_name),
            "read_mode": "header_and_type_sample_only",
            "data_imported": False,
        }
        if selected_sheet:
            inspection = self.inspector.inspect(source, sheet_name=None if source.suffix.lower() == ".csv" else selected_sheet, expected_headers=expected_headers)
            validation = self.inspector.validate_schema(inspection, expected_headers, allow_unexpected=False)
            result.update({
                "inspection": inspection,
                "validation": validation,
            })
        return result

    def import_file(
        self,
        path: Path,
        *,
        page_name: str,
        sheet_name: Optional[str] = None,
        period: Optional[str] = None,
        replace_current: bool = False,
    ) -> dict[str, Any]:
        """校验并导入文件；失败时整个批次回滚。"""
        source = Path(path)
        if not source.is_file():
            raise DataImportError(f"找不到导入文件: {source}")
        page, fields = self._load_page(page_name)
        primary_key = json.loads(page["primary_key_json"] or "[]")
        if len(primary_key) != 1:
            raise DataImportError("该数据表尚未确认唯一主键字段")
        primary_key = primary_key[0]
        confirmed_sheet = page["selected_sheet"]
        if confirmed_sheet and sheet_name != confirmed_sheet:
            raise DataImportError(f"必须选择已确认的工作表: {confirmed_sheet}")
        selected_sheet = sheet_name or confirmed_sheet
        # 真实业务数据必须与已确认 Schema 完全一致。此阶段不允许临时增加辅助列、
        # 重复列或未确认字段，避免错误文件被写入另一张结构相近的数据表。
        insert_fields = [field for field in fields if field["selected"]]
        expected_excel = [field["excel_name"] for field in insert_fields]
        inspection = self.inspector.inspect(
            source,
            sheet_name=selected_sheet,
            expected_headers=expected_excel,
        )
        validation = self.inspector.validate_schema(inspection, expected_excel, allow_unexpected=False)
        if not validation["ok"]:
            raise DataImportError(self._schema_error(validation))

        batch_id = uuid.uuid4().hex
        now = datetime.now(timezone.utc).isoformat()
        physical = self._physical_table(page_name)
        header_index = inspection["header_row"] - 1
        if header_index < 0:
            raise DataImportError("无法定位已识别的标题行")
        schema_hash = self.database.schema_fingerprint(page_name)
        columns = [field["technical_name"] for field in insert_fields]
        sql_columns = ["_sys_period", "_sys_batch_id", "_sys_imported_at", "_sys_source_file", *columns]
        placeholders = ", ".join("?" for _ in sql_columns)
        sql = f"INSERT INTO {self._quote_identifier(physical)} ({', '.join(self._quote_identifier(c) for c in sql_columns)}) VALUES ({placeholders})"
        imported = 0
        seen_primary_keys: set[str] = set()
        try:
            positions: dict[str, int] | None = None
            with self.database.connect_data() as conn:
                # 显式包住 DROP/CREATE 和后续写入，任何行转换失败都会恢复导入前的有效表。
                conn.execute("BEGIN IMMEDIATE")
                latest = conn.execute("SELECT schema_hash FROM import_batches WHERE page_name = ? AND status = 'completed' ORDER BY created_at DESC LIMIT 1", (page_name,)).fetchone()
                rebuild_table = replace_current or latest is None or not latest[0] or latest[0] != schema_hash
                physical = self.database.prepare_data_table(conn, page_name=page_name, fields=insert_fields, rebuild=rebuild_table)
                conn.execute("INSERT INTO import_batches(batch_id, page_name, period, source_file, record_count, status, schema_hash, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?)", (batch_id, page_name, period, source.name, 0, "running", schema_hash, now))
                # 只保留当前行并逐行写入 SQLite，避免把完整业务文件一次性载入内存。
                for row_number, row in enumerate(self._rows(source, selected_sheet), start=1):
                    if row_number == inspection["header_row"]:
                        header_names = self._normalize_headers(row)
                        positions = {name: index for index, name in enumerate(header_names)}
                        continue
                    if row_number <= inspection["header_row"]:
                        continue
                    if positions is None:
                        raise DataImportError("无法定位已识别的标题行")
                    if not any(value not in (None, "") for value in row):
                        continue
                    values = [self._sqlite_value(self._convert(row[positions[field["excel_name"]]] if positions[field["excel_name"]] < len(row) else None, field, row_number)) for field in insert_fields]
                    key_position = next((index for index, field in enumerate(insert_fields) if field["technical_name"] == primary_key), None)
                    if key_position is None or values[key_position] in (None, ""):
                        raise DataImportError(f"第 {row_number} 行主键字段 {primary_key} 不能为空")
                    key_value = str(values[key_position])
                    if key_value in seen_primary_keys:
                        raise DataImportError(f"第 {row_number} 行主键 {primary_key}={key_value} 重复")
                    seen_primary_keys.add(key_value)
                    conn.execute(sql, [period, batch_id, now, source.name, *values])
                    imported += 1
                conn.execute("UPDATE import_batches SET record_count = ?, status = ? WHERE batch_id = ?", (imported, "completed", batch_id))
            self.database.write_audit("data_imported", "data_page", page_name, {
                "batch_id": batch_id,
                "record_count": imported,
                "source_file": source.name,
                "sheet_name": selected_sheet,
                "schema_hash": schema_hash,
                "schema_validation": "exact_match",
            })
        except Exception as exc:
            self.database.write_audit("data_import_failed", "data_page", page_name, {"batch_id": batch_id, "source_file": source.name, "error": str(exc)})
            if isinstance(exc, DataImportError):
                raise
            raise DataImportError(f"第 {imported + 1} 条数据导入失败: {exc}") from exc
        return {
            "batch_id": batch_id,
            "page_name": page_name,
            "sheet_name": selected_sheet,
            "source_file": source.name,
            "record_count": imported,
            "schema_hash": schema_hash,
            "schema_validation": "exact_match",
            "status": "completed",
            "conversion_applied": True,
        }

    def _load_page(self, page_name: str) -> tuple[sqlite3.Row, list[dict[str, Any]]]:
        with self.database.connect_config() as conn:
            page = conn.execute("SELECT * FROM data_pages WHERE technical_name = ?", (page_name,)).fetchone()
            if page is None:
                raise DataImportError(f"未找到数据表: {page_name}")
            fields = [dict(row) for row in conn.execute("SELECT * FROM data_page_fields WHERE page_name = ? ORDER BY ordinal", (page_name,))]
            schema = conn.execute("SELECT definition_json FROM schema_versions WHERE page_name = ? ORDER BY confirmed_at DESC LIMIT 1", (page_name,)).fetchone()
            if schema is None:
                raise DataImportError("该数据表尚未确认表结构")
            definition = json.loads(schema[0])
            by_db = {item["technical_name"]: item for item in definition.get("columns", [])}
            for field in fields:
                source = by_db.get(field["technical_name"], {})
                field["excel_name"] = source.get("excel_name", field["display_name"])
                field["selected"] = bool(source.get("selected", True))
                field["transform_rule"] = source.get("transform_rule", "无需转换")
            return page, fields

    def _rows(self, path: Path, sheet_name: Optional[str]) -> Iterator[tuple[Any, ...]]:
        if path.suffix.lower() == ".csv":
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                yield from (tuple(row) for row in csv.reader(handle))
            return
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise DataImportError("读取 Excel 需要安装 openpyxl") from exc
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            name = sheet_name or book.sheetnames[0]
            if name not in book.sheetnames:
                raise DataImportError(f"未找到工作表: {name}")
            yield from book[name].iter_rows(values_only=True)
        finally:
            book.close()

    @staticmethod
    def _find_header(rows: Iterable[tuple[Any, ...]]) -> tuple[Any, ...]:
        for row in rows:
            values = [value for value in row if value not in (None, "")]
            if len(values) >= 2 and len({str(value).strip() for value in values}) == len(values):
                return row
        raise DataImportError("无法识别导入文件标题行")

    @staticmethod
    def _normalize_headers(row: tuple[Any, ...]) -> list[str]:
        seen: dict[str, int] = {}
        result = []
        for index, value in enumerate(row, start=1):
            name = str(value).strip() if value not in (None, "") else f"未命名列_{index}"
            seen[name] = seen.get(name, 0) + 1
            result.append(name if seen[name] == 1 else f"{name}_{seen[name]}")
        return result

    @staticmethod
    def _sqlite_value(value: Any) -> Any:
        # Decimal 不由 sqlite3 原生绑定；以十进制文本交给 SQLite NUMERIC/REAL 亲和性处理，避免在 Python 侧先转 float。
        return format(value, "f") if isinstance(value, Decimal) else value

    @staticmethod
    def _convert(value: Any, field: dict[str, Any], row_number: int) -> Any:
        if value in (None, ""):
            if field["required"]:
                raise DataImportError(f"第 {row_number} 行的字段 {field['technical_name']} 不能为空")
            return None
        target = field["sqlite_type"].upper()
        try:
            if target == "TEXT":
                return value.isoformat() if isinstance(value, (date, datetime)) else str(value)
            if target == "INTEGER":
                if isinstance(value, bool):
                    return int(value)
                return int(float(str(value).strip()))
            if target in {"REAL", "NUMERIC"}:
                parsed = Decimal(str(value).replace(",", "").strip())
                if not parsed.is_finite():
                    raise InvalidOperation("金额不是有限数")
                return parsed
        except (TypeError, ValueError, InvalidOperation) as exc:
            raise DataImportError(f"第 {row_number} 行的字段 {field['technical_name']} 无法转换为 {target}: {value}") from exc
        raise DataImportError(f"不支持的字段类型: {target}")

    @staticmethod
    def _schema_error(result: dict[str, Any]) -> str:
        parts = []
        if result["missing_columns"]:
            parts.append("缺少列：" + "、".join(result["missing_columns"]))
        if result["duplicated_columns"]:
            parts.append("重复列：" + "、".join(result["duplicated_columns"]))
        if result["unexpected_columns"]:
            parts.append("存在未确认列：" + "、".join(result["unexpected_columns"]))
        return "导入文件结构与已确认结构不一致；" + "；".join(parts)

    @staticmethod
    def _physical_table(page_name: str) -> str:
        if not re.fullmatch(r"[a-z][a-z0-9_]{0,62}", page_name):
            raise DataImportError("数据表技术名称不合法")
        return f"data_{page_name}"

    @staticmethod
    def _quote_identifier(identifier: str) -> str:
        if not re.fullmatch(r"[a-zA-Z_][a-zA-Z0-9_]*", identifier):
            raise DataImportError("数据库标识符不合法")
        return '"' + identifier + '"'
