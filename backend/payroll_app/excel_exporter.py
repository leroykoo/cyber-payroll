"""从 SQLite 结果表导出 Excel，并按输出方案写入大标题合并单元格。"""

from __future__ import annotations

import json
import sqlite3
from pathlib import Path
from typing import Any


class ExcelResultExporter:
    def export(self, data_connection: sqlite3.Connection, *, schema_connection: sqlite3.Connection, schema_name: str, output_path: Path) -> Path:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError as exc:
            raise RuntimeError("导出 Excel 需要安装 openpyxl") from exc
        schema = schema_connection.execute("SELECT fields_json, excel_headers_json FROM output_schemas WHERE schema_name = ?", (schema_name,)).fetchone()
        if schema is None:
            raise ValueError(f"未找到已保存的输出方案: {schema_name}")
        fields = json.loads(schema[0])
        layout = json.loads(schema[1])
        names = [field["name"] for field in fields]
        table = self._quote(schema_name)
        rows = data_connection.execute(f"SELECT {', '.join(self._quote(name) for name in names)} FROM {table}")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "本地核算结果"
        self._write_headers(sheet, fields, layout, Font, PatternFill, Alignment)
        for row_number, row in enumerate(rows, start=3):
            for column_number, value in enumerate(row, start=1):
                sheet.cell(row=row_number, column=column_number, value=value)
        sheet.freeze_panes = "A3"
        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output)
        return output

    @staticmethod
    def _write_headers(sheet: Any, fields: list[dict], layout: list[dict], Font: Any, PatternFill: Any, Alignment: Any) -> None:
        field_index = {field["name"]: index for index, field in enumerate(fields)}
        for index, field in enumerate(fields, start=1):
            sheet.cell(row=2, column=index, value=field.get("display_name", field["name"]))
        for cell in sheet[2]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="167C72")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for group in layout:
            if group.get("type") != "header":
                continue
            names = [name for name in group.get("field_names", []) if name in field_index]
            if not names:
                continue
            start = min(field_index[name] for name in names) + 1
            end = max(field_index[name] for name in names) + 1
            sheet.cell(row=1, column=start, value=group.get("title", ""))
            sheet.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
            cell = sheet.cell(row=1, column=start)
            cell.font = Font(bold=True, color="275A86")
            cell.fill = PatternFill("solid", fgColor="E8F2FC")
            cell.alignment = Alignment(horizontal="center", vertical="center")

    @staticmethod
    def _quote(identifier: str) -> str:
        if not identifier or not identifier[0].islower() or not all(char.isalnum() or char == "_" for char in identifier):
            raise ValueError(f"数据库标识符不合法: {identifier}")
        return '"' + identifier + '"'
