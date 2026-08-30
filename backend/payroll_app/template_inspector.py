"""以低内存方式检查 Excel/CSV 模板结构。

该模块只返回工作表元数据、标题行、列级类型统计和转换建议，不保存业务明细。
正式入库必须在用户确认 Schema 后由独立导入流程执行。
"""

from __future__ import annotations

import csv
import re
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable, Iterator, Optional


@dataclass
class ColumnSummary:
    ordinal: int
    excel_name: str
    non_empty: int
    type_counts: dict[str, int]
    excel_type: str
    suggested_sqlite_type: str
    suggested_transform_rule: str

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


class TemplateInspector:
    """流式读取模板，内存占用与文件总行数基本无关。"""

    def __init__(
        self,
        *,
        header_scan_rows: int = 20,
        type_sample_rows: int = 50,
        anomaly_sample_rows: int = 200,
    ) -> None:
        if header_scan_rows < 1:
            raise ValueError("header_scan_rows 必须大于 0")
        if type_sample_rows < 1:
            raise ValueError("type_sample_rows 必须大于 0")
        if anomaly_sample_rows < type_sample_rows:
            raise ValueError("anomaly_sample_rows 必须不小于 type_sample_rows")
        self.header_scan_rows = header_scan_rows
        self.type_sample_rows = type_sample_rows
        self.anomaly_sample_rows = anomaly_sample_rows

    def inspect(
        self,
        path: Path,
        *,
        sheet_name: Optional[str] = None,
        expected_headers: Optional[Iterable[str]] = None,
    ) -> dict[str, Any]:
        source = Path(path)
        suffix = source.suffix.lower()
        if suffix == ".csv":
            return self._inspect_rows(self._csv_rows(source), source.name, "CSV", None, expected_headers)
        if suffix in {".xlsx", ".xlsm"}:
            return self._inspect_xlsx(source, sheet_name, expected_headers)
        raise ValueError("仅支持 .xlsx、.xlsm 和 .csv 文件；旧版 .xls 需先转换为 .xlsx")

    @staticmethod
    def validate_schema(
        inspection: dict[str, Any],
        confirmed_columns: Iterable[str],
        *,
        allow_unexpected: bool = False,
    ) -> dict[str, Any]:
        """比较导入文件与已确认 Schema；只返回结果，不执行任何写入。

        正式导入只需要确认的映射列可用。Excel 中额外的列（包括重复表头
        被规范化后的 ``_2`` 列，以及末尾空白列）可以选择保留为提示信息而
        不阻断导入，避免模板存在无关辅助列时无法写入已确认的数据表。
        """
        actual = [column["excel_name"] for column in inspection.get("columns", [])]
        expected = list(confirmed_columns)
        actual_set = set(actual)
        expected_set = set(expected)
        missing = [name for name in expected if name not in actual_set]
        unexpected = [name for name in actual if name not in expected_set]
        duplicated = [name for name, count in Counter(actual).items() if count > 1]
        return {
            "ok": not missing and not duplicated and (allow_unexpected or not unexpected),
            "missing_columns": missing,
            "unexpected_columns": unexpected,
            "ignored_columns": unexpected if allow_unexpected else [],
            "duplicated_columns": duplicated,
            "header_row": inspection.get("header_row"),
            "sheet_name": inspection.get("sheet_name"),
        }

    def _inspect_xlsx(
        self,
        path: Path,
        sheet_name: Optional[str],
        expected_headers: Optional[Iterable[str]],
    ) -> dict[str, Any]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("读取 Excel 需要安装 openpyxl") from exc

        book = load_workbook(path, read_only=True, data_only=False)
        try:
            if sheet_name and sheet_name not in book.sheetnames:
                raise ValueError(f"未找到工作表: {sheet_name}")
            selected_sheet = sheet_name or book.sheetnames[0]
            rows = book[selected_sheet].iter_rows(values_only=True)
            return self._inspect_rows(rows, path.name, "Excel", selected_sheet, expected_headers)
        finally:
            book.close()

    @staticmethod
    def _csv_rows(path: Path) -> Iterator[tuple[Any, ...]]:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            for row in csv.reader(handle):
                yield tuple(row)

    def _inspect_rows(
        self,
        rows: Iterable[tuple[Any, ...]],
        source_file: str,
        file_type: str,
        sheet_name: Optional[str],
        expected_headers: Optional[Iterable[str]] = None,
    ) -> dict[str, Any]:
        iterator = iter(rows)
        preview = []
        for _ in range(self.header_scan_rows):
            try:
                preview.append(tuple(next(iterator)))
            except StopIteration:
                break
        if not preview:
            raise ValueError("文件没有可读取的内容")

        header_index = self._find_header_index(preview, expected_headers)
        headers = self._normalize_headers(preview[header_index])
        stats = [Counter() for _ in headers]
        record_count = 0
        sampled_rows = 0
        for row in preview[header_index + 1 :]:
            if sampled_rows >= self.type_sample_rows:
                break
            record_count += 1
            self._accumulate(stats, row, len(headers))
            sampled_rows += 1
        for row in iterator:
            if sampled_rows >= self.type_sample_rows:
                break
            record_count += 1
            self._accumulate(stats, row, len(headers))
            sampled_rows += 1

        # 常态只读少量样本；只有首轮已经发现混合类型或有效值过少时，才扩大本地复核范围。
        # 复核仍然不保存业务行，只继续维护列级类型计数。
        expanded_sample = False
        if sampled_rows >= self.type_sample_rows and self._needs_expanded_sample(stats):
            for row in iterator:
                if sampled_rows >= self.anomaly_sample_rows:
                    break
                record_count += 1
                self._accumulate(stats, row, len(headers))
                sampled_rows += 1
            expanded_sample = sampled_rows > self.type_sample_rows

        columns = []
        for index, (name, counts) in enumerate(zip(headers, stats), start=1):
            columns.append(self._summarize_column(index, name, counts))
        return {
            "stage": "schema_preview",
            "source_file": source_file,
            "file_type": file_type,
            "sheet_name": sheet_name,
            "header_row": header_index + 1,
            "record_count": None,
            "sample_record_count": record_count,
            "sample_limit": self.type_sample_rows,
            "sample_max_limit": self.anomaly_sample_rows,
            "sample_expanded": expanded_sample,
            "read_mode": "header_and_type_sample_only",
            "data_imported": False,
            "conversion_applied": False,
            "columns": [column.to_dict() for column in columns],
            "memory_mode": "streaming_header_and_bounded_sample",
        }

    @classmethod
    def _needs_expanded_sample(cls, stats: list[Counter]) -> bool:
        if any(cls._summarize_column(index, "", counts).excel_type == "混合类型" for index, counts in enumerate(stats, start=1)):
            return True
        return any(sum(count for kind, count in counts.items() if kind != "empty") < 3 for counts in stats)

    @staticmethod
    def _find_header_index(
        rows: list[tuple[Any, ...]],
        expected_headers: Optional[Iterable[str]] = None,
    ) -> int:
        expected = {str(value).strip() for value in (expected_headers or []) if str(value).strip()}

        def values(row: tuple[Any, ...]) -> list[str]:
            return [str(value).strip() for value in row if value not in (None, "")]

        # 已确认结构优先：真实表头通常位于说明行之后，不能仅按“非空单元格最多”判断。
        # 只有完全找不到已确认字段时，才回退到原来的密度判断，兼容新建模板预览。
        if expected:
            matched = []
            for index, row in enumerate(rows):
                row_values = values(row)
                overlap = len(set(row_values) & expected)
                if overlap:
                    matched.append((index, overlap, len(row_values), len(set(row_values))))
            if matched:
                return max(matched, key=lambda item: (item[1], item[2], item[3], -item[0]))[0]

        def score(row: tuple[Any, ...]) -> tuple[int, int]:
            row_values = values(row)
            unique = len(set(row_values))
            return (len(row_values), unique)

        candidates = [(index, score(row)) for index, row in enumerate(rows) if score(row)[0] >= 2]
        if not candidates:
            raise ValueError("无法在前置扫描范围内识别标题行")
        return max(candidates, key=lambda item: (item[1][0], item[1][1], -item[0]))[0]

    @staticmethod
    def _normalize_headers(row: tuple[Any, ...]) -> list[str]:
        headers = []
        seen: Counter[str] = Counter()
        for index, value in enumerate(row, start=1):
            raw = str(value).strip() if value not in (None, "") else f"未命名列_{index}"
            seen[raw] += 1
            headers.append(raw if seen[raw] == 1 else f"{raw}_{seen[raw]}")
        return headers

    @staticmethod
    def _accumulate(stats: list[Counter], row: tuple[Any, ...], width: int) -> None:
        for index in range(width):
            value = row[index] if index < len(row) else None
            kind = TemplateInspector._cell_kind(value)
            stats[index][kind] += 1

    @staticmethod
    def _cell_kind(value: Any) -> str:
        if value in (None, ""):
            return "empty"
        if isinstance(value, bool):
            return "boolean"
        if isinstance(value, (datetime, date)):
            return "date"
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return "number"
        if isinstance(value, str) and value.startswith("="):
            return "formula"
        if isinstance(value, str) and re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)", value.strip()):
            return "numeric_text"
        return "text"

    @staticmethod
    def _summarize_column(ordinal: int, name: str, counts: Counter) -> ColumnSummary:
        non_empty = sum(count for kind, count in counts.items() if kind != "empty")
        observed = {kind for kind, count in counts.items() if count and kind != "empty"}
        numeric_like = observed and observed <= {"number", "numeric_text"}
        if not observed:
            excel_type, sqlite_type, transform = "空列", "TEXT", "待确认"
        elif observed <= {"date"}:
            excel_type, sqlite_type, transform = "日期", "TEXT", "日期标准化"
        elif numeric_like:
            excel_type, sqlite_type, transform = "数值" if "number" in observed else "文本数字", "REAL", "文本转数值" if "numeric_text" in observed else "无需转换"
        elif observed <= {"boolean"}:
            excel_type, sqlite_type, transform = "布尔", "INTEGER", "布尔转整数"
        elif "number" in observed or "numeric_text" in observed or "date" in observed:
            excel_type, sqlite_type, transform = "混合类型", "TEXT", "需人工确认"
        else:
            excel_type, sqlite_type, transform = "文本", "TEXT", "无需转换"
        return ColumnSummary(ordinal, name, non_empty, dict(counts), excel_type, sqlite_type, transform)
